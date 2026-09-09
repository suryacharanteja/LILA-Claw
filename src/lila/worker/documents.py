"""Local extraction only. Imported facts remain proposals for owner review."""
import io
import re
from lila.domain.artifacts import validate_document

MAX_CHARACTERS=100000
LABELS={'name':'name','email':'email','phone':'phone','location':'location',
        'work authorization':'work_authorization','notice period':'notice_period'}


def extract(data,media_type):
    validate_document(data,media_type)
    parts=[]
    count=0
    if media_type=='application/pdf':
        from pypdf import PdfReader
        document=PdfReader(io.BytesIO(data),strict=True)
        iterator=(page.extract_text() or '' for page in document.pages)
    else:
        from docx import Document
        document=Document(io.BytesIO(data))
        def paragraphs():
            for paragraph in document.paragraphs:
                yield paragraph.text
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield cell.text
        iterator=paragraphs()
    truncated=False
    for part in iterator:
        part=part.replace('\x00','')
        available=MAX_CHARACTERS-count
        if len(part)+1>available:
            parts.append(part[:max(0,available)])
            truncated=True
            break
        parts.append(part)
        count+=len(part)+1
    text='\n'.join(parts)[:MAX_CHARACTERS].strip()
    proposals=[]
    seen=set()
    for line in text.splitlines():
        match=re.fullmatch(r'\s*([A-Za-z ]{1,40}):\s*(.{1,1000})\s*',line)
        if not match:
            continue
        key=LABELS.get(match[1].strip().lower())
        if key and (key,match[2]) not in seen:
            seen.add((key,match[2]))
            proposals.append({'field_key':key,'value':match[2].strip()})
    return {'text':text,'proposals':proposals[:100],
            'status':'MANUAL_ENTRY_REQUIRED' if not text else 'TRUNCATED' if truncated else 'EXTRACTED'}
