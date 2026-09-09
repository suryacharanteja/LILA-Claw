import asyncio
import io
import ssl
import time
from uuid import uuid4
from docx import Document
from lila.runtime.installation import prepare
from lila.runtime.service import Runtime
from lila.runtime.worker_http import async_client


def test_real_worker_extracts_proposals_without_overwriting_owner(tmp_path):
    root=tmp_path/'installation'
    prepare(root)
    runtime=Runtime(root,require_trust=False,preferred_port=0,embedded_test_worker=True)
    try:
        assert runtime.start()
        domain=runtime.domain
        owner=str(uuid4())
        account=domain.register_account('extraction-fixture')
        domain.fact_command(owner,account,{'command_id':str(uuid4()),'operation':'verify','field_key':'name','value':'Owner confirmed name'})
        document=Document()
        document.add_paragraph('Name: Imported proposed name')
        document.add_paragraph('Email: fixture@example.invalid')
        output=io.BytesIO()
        document.save(output)
        media='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        upload=domain.documents.accept(owner,str(uuid4()),account,'fixture.docx',output.getvalue(),media)
        domain.documents.publish(upload['upload_id'],output.getvalue(),media)
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            row=domain.writer.call(lambda db:db.execute("SELECT phase FROM operations WHERE kind='DOCUMENT_EXTRACT'").fetchone())
            if row==('COMPLETE',):
                break
            time.sleep(0.1)
        assert row==('COMPLETE',)
        facts=domain.writer.call(lambda db:db.execute('SELECT f.field_key,v.status,c.canonical_json FROM facts f JOIN fact_versions v ON v.id=f.current_version JOIN content_versions c ON c.id=v.value_version ORDER BY f.field_key').fetchall())
        assert facts==[('email','PROPOSED','"fixture@example.invalid"'),('name','VERIFIED','"Owner confirmed name"')]
        assert domain.writer.call(lambda db:db.execute('SELECT count(*) FROM fact_versions').fetchone())[0]==3
        assert runtime.workers.healthy()
        assert runtime.worker.process.pid!=__import__('os').getpid()
        assert domain.verify_audit()['events_verified']>0
    finally:
        runtime.close()


def test_async_worker_pin_mismatch_rejected(tmp_path):
    root=tmp_path/'installation'
    prepare(root)
    runtime=Runtime(root,require_trust=False,preferred_port=0)
    try:
        assert runtime.start()
        boot=runtime.control('worker_boot')
        boot['server_spki_pin']='0'*64
        async def probe():
            import pytest
            async with async_client(boot) as client:
                with pytest.raises(ssl.SSLError):
                    await client.post('/internal/v1/health',json={})
        asyncio.run(probe())
    finally:
        runtime.close()
